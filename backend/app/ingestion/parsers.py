"""Document parsers (CLAUDE.md §6.4).

pdf=PyMuPDF, docx=python-docx, html=BeautifulSoup, txt/md=plain. Each returns a
normalized text plus a page map (char offsets per page) for span tracing, and a
sha256 content hash over the normalized text. Heavy parsers are lazy-imported so
this module loads even when a given parser lib is absent.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from typing import Any

from app.core.exceptions import ValidationAppError


@dataclass
class PageSpan:
    page_number: int
    char_start: int
    char_end: int


@dataclass
class ParsedDoc:
    text: str
    pages: list[PageSpan]
    content_hash: str
    meta: dict[str, Any] = field(default_factory=dict)


def _normalize(text: str) -> str:
    """Normalize newlines so offsets are stable across platforms."""
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _single_page(text: str) -> list[PageSpan]:
    return [PageSpan(page_number=1, char_start=0, char_end=len(text))]


def _parse_text(path: str) -> ParsedDoc:
    with open(path, encoding="utf-8", errors="replace") as fh:
        text = _normalize(fh.read())
    return ParsedDoc(text=text, pages=_single_page(text), content_hash=_hash(text))


def _parse_html(path: str) -> ParsedDoc:
    from bs4 import BeautifulSoup  # lazy

    with open(path, encoding="utf-8", errors="replace") as fh:
        soup = BeautifulSoup(fh.read(), "html.parser")
    text = _normalize(soup.get_text(separator="\n"))
    return ParsedDoc(text=text, pages=_single_page(text), content_hash=_hash(text))


def _parse_docx(path: str) -> ParsedDoc:
    import docx  # lazy (python-docx)

    document = docx.Document(path)
    text = _normalize("\n".join(p.text for p in document.paragraphs))
    return ParsedDoc(
        text=text,
        pages=_single_page(text),
        content_hash=_hash(text),
        meta={"paragraphs": len(document.paragraphs)},
    )


def _parse_pdf(path: str) -> ParsedDoc:
    import fitz  # lazy (PyMuPDF)

    parts: list[str] = []
    pages: list[PageSpan] = []
    pos = 0
    with fitz.open(path) as doc:
        for i, page in enumerate(doc):
            page_text = _normalize(page.get_text())
            start = pos
            parts.append(page_text)
            pos += len(page_text)
            pages.append(PageSpan(page_number=i + 1, char_start=start, char_end=pos))
            # Page separator (belongs to no page span).
            parts.append("\n")
            pos += 1
    text = "".join(parts)
    return ParsedDoc(
        text=text, pages=pages, content_hash=_hash(text), meta={"page_count": len(pages)}
    )


_PARSERS = {
    "txt": _parse_text,
    "md": _parse_text,
    "html": _parse_html,
    "docx": _parse_docx,
    "pdf": _parse_pdf,
}


def parse_document(path: str, filetype: str) -> ParsedDoc:
    parser = _PARSERS.get(filetype.lower())
    if parser is None:
        raise ValidationAppError(f"Unsupported file type: {filetype}")
    return parser(path)
